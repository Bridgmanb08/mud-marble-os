import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from .. import line_items
from ..change_order_utils import compute_sop_breach
from ..deps import CurrentUser, get_current_user
from ..pdf_export import (
    STATUS_COLORS,
    NumberedCanvas,
    SIDE_MARGIN,
    breadcrumb_for,
    build_info_card,
    build_line_items_table,
    build_letterhead,
    build_styles,
    build_totals_band,
    xml_escape,
)
from ..schemas.change_orders import ChangeOrderCreate, ChangeOrderOut, ChangeOrderUpdate
from ..schemas.estimates import LineItemCreate, LineItemOut, LineItemUpdate
from ..supabase_client import db_delete, db_get, db_patch, db_post

router = APIRouter(prefix="/change-orders", tags=["change_orders"])

# Shared between the PDF and Excel exports below (and previously duplicated
# inline in just the PDF export) -- one place to add a new co_type/
# discovered_by option's display label instead of two.
TYPE_LABELS = {"client_addition": "Client Addition", "oversight": "Oversight", "unforeseen": "Unforeseen"}
DISCOVERED_LABELS = {"brent": "Brent", "shannon": "Shannon", "client": "Client", "subcontractor": "Subcontractor"}


def _attach_breach(co: dict) -> dict:
    now = datetime.now(timezone.utc)
    return {**co, "sop_breach": compute_sop_breach(co.get("status"), co.get("sent_at"), now)}


async def _recalc_co_totals(co_id: str) -> None:
    """Once a change order has real line items, its owner_price/builder_cost
    stop being something typed directly onto the CO and become the sum of
    those items instead -- the same "the parent's totals are always
    recomputed from its line items, never independently trusted" rule
    estimates.py's own _recalc_estimate_totals already follows. Only ever
    called from the three line-item endpoints below (create/update/delete),
    never from anywhere that would touch a CO that's never had a line item
    added -- so the original flat-entry flow (manually typed owner price,
    no line items at all) is never affected by this running. Deliberately
    always persists the sum, including 0 once the last item is removed,
    rather than leaving a stale nonzero total behind."""
    items = await db_get("estimate_line_items", f"?change_order_id=eq.{co_id}&select=builder_cost,owner_price")
    total_builder_cost = round(sum(i.get("builder_cost") or 0 for i in items), 2)
    total_owner_price = round(sum(i.get("owner_price") or 0 for i in items), 2)

    co_rows = await db_get("change_orders", f"?id=eq.{co_id}&select=status,project_id,owner_price")
    if not co_rows:
        return
    old_owner_price = co_rows[0].get("owner_price") or 0
    project_id = co_rows[0]["project_id"]
    status = co_rows[0]["status"]

    await db_patch("change_orders", co_id, {"builder_cost": total_builder_cost, "owner_price": total_owner_price})

    # An already-approved CO's owner_price is already baked into the
    # project's contract_value (see update_change_order's status-transition
    # handling) -- a line-item edit that changes the total while status
    # stays "approved" throughout has to adjust contract_value by the
    # delta too, or an approved CO's price could silently drift away from
    # what the project's own contract total shows.
    if status == "approved" and total_owner_price != old_owner_price:
        delta = total_owner_price - old_owner_price
        proj_rows = await db_get("projects", f"?id=eq.{project_id}&select=contract_value")
        if proj_rows:
            current = proj_rows[0].get("contract_value") or 0
            await db_patch("projects", project_id, {"contract_value": round(current + delta, 2)})


@router.get("", response_model=list[ChangeOrderOut])
async def list_change_orders(project_id: Optional[str] = None, _: CurrentUser = Depends(get_current_user)):
    query = "?order=created_at.desc&select=*,projects(name,address)"
    if project_id:
        query += f"&project_id=eq.{project_id}"
    rows = await db_get("change_orders", query)
    return [_attach_breach(r) for r in rows]


@router.post("", response_model=ChangeOrderOut)
async def create_change_order(body: ChangeOrderCreate, _: CurrentUser = Depends(get_current_user)):
    existing = await db_get(
        "change_orders", f"?project_id=eq.{body.project_id}&select=co_number&order=co_number.desc&limit=1"
    )
    next_number = (existing[0]["co_number"] + 1) if existing and existing[0].get("co_number") else 1
    data = body.model_dump(exclude_none=True)
    data["co_number"] = next_number
    data["status"] = "pending"
    data["sent_at"] = datetime.now(timezone.utc).isoformat()
    rows = await db_post("change_orders", data)
    full = await db_get("change_orders", f"?id=eq.{rows[0]['id']}&select=*,projects(name,address)")
    return _attach_breach(full[0])


@router.patch("/{co_id}", response_model=ChangeOrderOut)
async def update_change_order(co_id: str, body: ChangeOrderUpdate, _: CurrentUser = Depends(get_current_user)):
    existing = await db_get("change_orders", f"?id=eq.{co_id}&select=status,project_id,owner_price")
    if not existing:
        raise HTTPException(status_code=404, detail="Change order not found")
    old_status = existing[0]["status"]
    project_id = existing[0]["project_id"]
    owner_price = existing[0].get("owner_price") or 0

    # exclude_unset (not exclude_none) -- a caller may need to explicitly clear
    # a field (e.g. an import correction clearing description), and that null
    # has to reach the database instead of being silently dropped.
    await db_patch("change_orders", co_id, body.model_dump(exclude_unset=True))

    # Approving a change order should grow the project's contracted value by
    # its owner price (and shrink it back if it's ever un-approved) -- without
    # this, an approved addition never shows up in the project's contract
    # value or the dashboard's total-contract-value rollup.
    new_status = body.status if body.status is not None else old_status
    if new_status != old_status and owner_price:
        delta = 0.0
        if new_status == "approved" and old_status != "approved":
            delta = owner_price
        elif old_status == "approved" and new_status != "approved":
            delta = -owner_price
        if delta:
            proj_rows = await db_get("projects", f"?id=eq.{project_id}&select=contract_value")
            if proj_rows:
                current = proj_rows[0].get("contract_value") or 0
                await db_patch("projects", project_id, {"contract_value": round(current + delta, 2)})

    full = await db_get("change_orders", f"?id=eq.{co_id}&select=*,projects(name,address)")
    return _attach_breach(full[0])


@router.delete("/{co_id}")
async def delete_change_order(co_id: str, _: CurrentUser = Depends(get_current_user)):
    existing = await db_get("change_orders", f"?id=eq.{co_id}&select=status")
    if not existing:
        raise HTTPException(status_code=404, detail="Change order not found")
    # An approved CO already grew the project's contract_value (see
    # update_change_order above) -- the same "don't let a hard delete
    # destroy real agreed-upon financial history" guard invoices.py applies
    # to a paid invoice, here applied to an approved change order instead
    # of trying to reverse that contract_value math on delete.
    if existing[0]["status"] == "approved":
        raise HTTPException(
            status_code=400,
            detail="Can't delete an approved change order -- change its status first if it was approved in error.",
        )
    await db_delete("change_orders", co_id)
    return {"ok": True}


@router.get("/{co_id}", response_model=ChangeOrderOut)
async def get_change_order(co_id: str, _: CurrentUser = Depends(get_current_user)):
    rows = await db_get("change_orders", f"?id=eq.{co_id}&select=*,projects(name,address)")
    if not rows:
        raise HTTPException(status_code=404, detail="Change order not found")
    return _attach_breach(rows[0])


@router.get("/{co_id}/items", response_model=list[LineItemOut])
async def list_co_items(co_id: str, _: CurrentUser = Depends(get_current_user)):
    return await line_items.list_items("change_order_id", co_id)


@router.post("/{co_id}/items", response_model=LineItemOut)
async def create_co_item(co_id: str, body: LineItemCreate, _: CurrentUser = Depends(get_current_user)):
    existing = await db_get("change_orders", f"?id=eq.{co_id}&select=id")
    if not existing:
        raise HTTPException(status_code=404, detail="Change order not found")
    item = await line_items.create_item("change_order_id", co_id, body)
    await _recalc_co_totals(co_id)
    return item


@router.patch("/{co_id}/items/{item_id}", response_model=LineItemOut)
async def update_co_item(co_id: str, item_id: str, body: LineItemUpdate, _: CurrentUser = Depends(get_current_user)):
    item = await line_items.update_item(item_id, body)
    await _recalc_co_totals(co_id)
    return item


@router.delete("/{co_id}/items/{item_id}")
async def delete_co_item(co_id: str, item_id: str, _: CurrentUser = Depends(get_current_user)):
    await line_items.delete_item(item_id)
    await _recalc_co_totals(co_id)
    return {"ok": True}


@router.get("/{co_id}/export/pdf")
async def export_change_order_pdf(co_id: str, _: CurrentUser = Depends(get_current_user)):
    rows = await db_get("change_orders", f"?id=eq.{co_id}&select=*,projects(name,address,clients(first_name,last_name))")
    if not rows:
        raise HTTPException(status_code=404, detail="Change order not found")
    co = _attach_breach(rows[0])
    project = co.get("projects") or {}
    breadcrumb = breadcrumb_for(project)
    project_name = (project.get("name") or "").split("|")[0].strip()
    items = await line_items.list_items("change_order_id", co_id)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter, topMargin=0.5 * inch, bottomMargin=0.5 * inch, leftMargin=SIDE_MARGIN, rightMargin=SIDE_MARGIN
    )
    PAGE_WIDTH = letter[0] - 2 * SIDE_MARGIN
    s = build_styles()

    elements = build_letterhead(s, PAGE_WIDTH, breadcrumb)
    co_number = f"CO-{str(co.get('co_number') or '?').zfill(3)}"
    elements.append(Paragraph(f"Change Order {co_number}: {xml_escape(co.get('title') or '')}", s["title"]))
    address = (project.get("address") or "").strip() or project_name
    if address:
        elements.append(Paragraph(xml_escape(address), s["address"]))
    elements.append(HRFlowable(width="100%", thickness=0.75, color=colors.lightgrey, spaceAfter=12))

    # A bordered info card, same treatment as the invoice PDF -- status
    # colored to match its on-screen badge (STATUS_BADGE in
    # ChangeOrders.tsx) so "approved" reads as unmistakably good news.
    status_color_key = {"approved": "green", "rejected": "red", "sent": "amber", "pending": "gray"}.get(co.get("status"))
    info_fields = [
        ("Type", xml_escape(TYPE_LABELS.get(co.get("co_type"), co.get("co_type") or "")), None),
        ("Status", xml_escape((co.get("status") or "").title()), STATUS_COLORS.get(status_color_key)),
        ("Discovered by", xml_escape(DISCOVERED_LABELS.get(co.get("discovered_by"), co.get("discovered_by") or "—")), None),
    ]
    elements.append(build_info_card(s, PAGE_WIDTH, info_fields, columns=3))
    elements.append(Spacer(1, 14))

    # Client-facing scope description only -- notes_internal is exactly the
    # team-only field it's named for and must never reach a document a
    # client could see, same principle as unit_cost/builder_cost being
    # excluded from the estimate export.
    if co.get("description"):
        elements.append(Paragraph(xml_escape(co["description"]), s["body"]))
        elements.append(Spacer(1, 10))

    # Once a change order has real line items, the flat description above is
    # no longer where the scope detail lives -- without this table the PDF
    # only ever showed a single lump "Price" line, which is exactly what
    # made a line-itemized CO read as vague. Same Item/Description/Qty·Unit/
    # Unit Price/Price table the estimate PDF uses (build_line_items_table),
    # so a line item's scope reads identically on both documents; cost codes
    # and builder_cost stay off this client-facing table, same rule the
    # estimate export follows.
    if items:
        item_rows = []
        for i in items:
            qty = i.get("quantity") or 0
            owner_price = i.get("owner_price") or 0
            unit_price = (owner_price / qty) if qty else owner_price
            item_rows.append((i.get("title"), line_items.client_text(i), qty, i.get("unit"), unit_price, owner_price))
        elements.append(build_line_items_table(s, PAGE_WIDTH, item_rows))
        elements.append(Spacer(1, 12))

    # Owner price only -- builder_cost is internal margin data, same rule
    # the estimate PDF already follows for its own owner-price/builder-cost
    # split. Same totals-band treatment as the invoice PDF's Amount Due.
    owner_price = co.get("owner_price") or 0
    elements.append(build_totals_band(s, PAGE_WIDTH, [("Price", f"${owner_price:,.2f}", True)]))

    # A change order needs the client's sign-off to actually be approved --
    # unlike an invoice (just a bill), this is an agreement, so it gets the
    # same signature block the estimate proposal PDF ends with. The one-line
    # instruction above it is what the estimate's own "A signed copy
    # authorizes work to begin" line already does -- without it, a client
    # reaches three blank signature lines with no explanation of what
    # signing means.
    elements.append(Spacer(1, 20))
    elements.append(Paragraph("Please sign below to approve this change order.", s["body"]))
    elements.append(Spacer(1, 16))
    elements.append(Paragraph("Signature: _______________________________________", s["body"]))
    elements.append(Spacer(1, 10))
    elements.append(Paragraph("Date: _______________________________________", s["body"]))
    elements.append(Spacer(1, 10))
    elements.append(Paragraph("Print Name: _______________________________________", s["body"]))

    doc.build(elements, canvasmaker=NumberedCanvas)
    pdf_bytes = buf.getvalue()
    buf.close()

    filename = f"{co_number}-{project_name or 'change-order'}.pdf".replace(" ", "-")
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{co_id}/export/excel")
async def export_change_order_excel(co_id: str, _: CurrentUser = Depends(get_current_user)):
    rows = await db_get("change_orders", f"?id=eq.{co_id}&select=*,projects(name,address)")
    if not rows:
        raise HTTPException(status_code=404, detail="Change order not found")
    co = _attach_breach(rows[0])
    project = co.get("projects") or {}
    project_name = (project.get("name") or "").split("|")[0].strip()
    co_number = f"CO-{str(co.get('co_number') or '?').zfill(3)}"
    items = await line_items.list_items("change_order_id", co_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Change Order"
    header_font = Font(bold=True)

    ws.append([f"Change Order {co_number}: {co.get('title') or ''}"])
    ws["A1"].font = Font(bold=True, size=14)
    address = (project.get("address") or "").strip() or project_name
    if address:
        ws.append([address])
    ws.append([])

    ws.append(["Type", "Status", "Discovered by"])
    for cell in ws[ws.max_row]:
        cell.font = header_font
    ws.append(
        [
            TYPE_LABELS.get(co.get("co_type"), co.get("co_type") or ""),
            (co.get("status") or "").title(),
            DISCOVERED_LABELS.get(co.get("discovered_by"), co.get("discovered_by") or "—"),
        ]
    )
    ws.append([])

    # Client-facing scope description only -- notes_internal never belongs
    # in anything a client could receive, same rule the PDF export follows.
    if co.get("description"):
        ws.append(["Description"])
        ws.cell(row=ws.max_row, column=1).font = header_font
        ws.append([co["description"]])
        ws.append([])

    # Once a change order has real line items, list them with the exact same
    # columns as the estimate Excel export (Item/Description/Qty/Unit/Unit
    # Price/Price) -- otherwise this sheet has the same "flat lump price, no
    # scope detail" gap the PDF export had, and read differently from an
    # estimate's own line items for no real reason.
    if items:
        ws.append(["Item", "Description", "Qty", "Unit", "Unit Price", "Price"])
        for cell in ws[ws.max_row]:
            cell.font = header_font
        for it in items:
            qty = it.get("quantity") or 0
            owner_price = it.get("owner_price") or 0
            unit_price = (owner_price / qty) if qty else owner_price
            ws.append([it.get("title"), line_items.client_text(it), it.get("quantity"), it.get("unit"), unit_price, owner_price])
        ws.append([])

    ws.append(["", "", "", "", "Price", co.get("owner_price") or 0])
    ws.cell(row=ws.max_row, column=5).font = header_font
    ws.cell(row=ws.max_row, column=6).font = header_font

    for col, width in zip("ABCDEF", [28, 40, 8, 8, 12, 12]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()
    buf.close()

    filename = f"{co_number}-{project_name or 'change-order'}.xlsx".replace(" ", "-")
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
